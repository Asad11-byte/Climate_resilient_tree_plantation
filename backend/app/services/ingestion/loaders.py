"""
File loaders: raw bytes on disk -> plain text (+ per-page text for PDFs).
No chunking or embedding here — this module has exactly one job: extraction.
"""
import hashlib
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional

import docx
from pypdf import PdfReader

from app.core.exceptions import AppError

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class UnsupportedFileTypeError(AppError):
    pass


@dataclass
class RawDocument:
    document_id: str          # stable hash of file content — same file re-ingested keeps the same id
    title: str
    source_path: str
    file_type: str
    text: str                          # full plain text (used for txt/md/docx)
    pages: Optional[List[str]] = None  # per-page text (used for pdf); None otherwise
    extra: dict = field(default_factory=dict)


def _content_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()[:16]


def _load_pdf(path: Path) -> RawDocument:
    data = path.read_bytes()
    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return RawDocument(
        document_id=_content_hash(data),
        title=path.stem,
        source_path=str(path),
        file_type="pdf",
        text="\n\n".join(pages),
        pages=pages,
    )


def _load_docx(path: Path) -> RawDocument:
    data = path.read_bytes()
    document = docx.Document(str(path))
    parts = []
    for para in document.paragraphs:
        if not para.text.strip():
            continue
        style = (para.style.name or "").lower() if para.style else ""
        if style.startswith("heading"):
            level = "#" * min(int(style[-1]), 6) if style[-1].isdigit() else "##"
            parts.append(f"{level} {para.text.strip()}")
        else:
            parts.append(para.text.strip())
    return RawDocument(
        document_id=_content_hash(data),
        title=path.stem,
        source_path=str(path),
        file_type="docx",
        text="\n\n".join(parts),
    )


def _load_text(path: Path, file_type: str) -> RawDocument:
    data = path.read_bytes()
    return RawDocument(
        document_id=_content_hash(data),
        title=path.stem,
        source_path=str(path),
        file_type=file_type,
        text=data.decode("utf-8", errors="replace"),
    )


def load_file(path: Path) -> RawDocument:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix == ".docx":
        return _load_docx(path)
    if suffix == ".txt":
        return _load_text(path, "txt")
    if suffix == ".md":
        return _load_text(path, "md")
    raise UnsupportedFileTypeError(
        f"Unsupported file type '{suffix}'. Supported: {sorted(SUPPORTED_EXTENSIONS)}"
    )
